"""
Extracts raw text from Word (.docx) files using python-docx.
Covers both regular paragraphs and text inside tables, since company
docs (FAQs, policy sheets) often put content in tables.
"""
from docx import Document


def parse_docx(file_path: str) -> str:
    """Extracts and concatenates text from paragraphs and tables in a .docx file."""
    doc = Document(file_path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)
